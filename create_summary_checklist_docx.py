import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="B0BEC5", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def build_summary_checklist_doc(output_path):
    doc = docx.Document()
    
    # 1.0 inch page margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
    # Document Header Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("SRH LINK-TODA APPLICATION")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Appendix C: User Manual — Summary Checklist of Figures & Functional Modules")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    # Introductory Narrative
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    r_intro = p_intro.add_run(
        "The following master summary checklist indexes all twenty-one (21) user manual figures documented across the eleven (11) functional modules of the SRH LINK-TODA progressive web application (Appendix C). Each entry outlines the assigned figure number, corresponding document section and UI form/interface name, primary target user roles, and an exhaustive technical description of the core features, interactive controls, and backend workflows covered."
    )
    r_intro.font.name = "Times New Roman"
    r_intro.font.size = Pt(10.5)
    r_intro.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    # Notes Section
    p_notes_hdr = doc.add_paragraph()
    p_notes_hdr.paragraph_format.space_before = Pt(8)
    p_notes_hdr.paragraph_format.space_after = Pt(4)
    p_notes_hdr.paragraph_format.keep_with_next = True
    r_nh = p_notes_hdr.add_run("Notes:")
    r_nh.font.name = "Times New Roman"
    r_nh.font.size = Pt(11)
    r_nh.font.bold = True
    r_nh.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    notes_items = [
        "The screenshots presented in this User Manual were captured from the current implemented version of the SRH LINK-TODA: A Progressive Web Application–Based Queue Management System with Geolocation, Geofencing, and Priority-Based Queue Algorithm during the capstone evaluation period. Interface appearance, active UI themes, and available controls may vary depending on the authenticated user role, account accreditation status, device viewport, and dynamic operational state (e.g., in-queue, on-trip, or en route).",
        "All names, passenger contact details, driver profiles, municipal MTOP (Motorized Tricycle Operator's Permit) franchise numbers, vehicle plate numbers, GPS coordinates, trip routes, transaction references, fare amounts, dispute narratives, and official announcements shown in the screenshots are fictional or test data used solely for academic evaluation and documentation purposes. They do not represent actual commuters, registered TODA operators, real-world passenger records, or official association records.",
        "SRH LINK-TODA applies role-based access controls. Unregistered public visitors, registered commuters, accredited TODA drivers, station dispatchers, and system superadministrators are provided only with the functions and information permitted for their respective roles. Confidential information—such as private passenger contact details, live driver queue position telemetry, scanned franchise documentation, earnings ledgers, and formal grievance proceedings—is protected against unauthorized access or public disclosure.",
        "Some pages contain additional controls or information further down the same screen. To maintain readability and avoid unnecessary repetition, this manual presents representative screenshots and focuses on the controls, state transitions, validation mechanisms, and procedures most relevant to completing each user task.",
        "Status labels, queue rankings, classification information, distance tariff computations, geofence check-in verifications, analytics, and other system-generated summaries are intended to support fair dispatching, monitoring, documentation, and authorized administrative review. They should not be interpreted as automatic legal determinations or replacements for official municipal ordinances, Barangay regulatory guidelines, or the supervisory assessment of presiding TODA dispatch officers.",
        "SRH LINK-TODA is an intelligent transit queue dispatch and fare monitoring application; it is not an emergency roadside rescue or public safety dispatch service. When a person is in immediate danger, requires urgent roadside assistance, or encounters vehicular accidents or safety threats en route, the appropriate emergency authorities, police (PNP), Barangay Tanods, or Cabanatuan City emergency response units should be contacted through the proper official channels.",
        "This User Manual reflects the system version documented at the time of the capstone study. Minor interface adjustments may be introduced during subsequent refinement while retaining the established system functions, priority-based queuing algorithm, geofencing parameters, privacy controls, and role-based workflows."
    ]
    
    for idx, item in enumerate(notes_items, 1):
        p_note = doc.add_paragraph()
        p_note.paragraph_format.space_before = Pt(2)
        p_note.paragraph_format.space_after = Pt(3)
        p_note.paragraph_format.line_spacing = 1.15
        p_note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        r_num = p_note.add_run(f"{idx}. ")
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(10)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        
        r_txt = p_note.add_run(item)
        r_txt.font.name = "Times New Roman"
        r_txt.font.size = Pt(10)
        r_txt.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    # Table Caption
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(12)
    p_cap.paragraph_format.space_after = Pt(4)
    p_cap.paragraph_format.keep_with_next = True
    
    r_cap_label = p_cap.add_run("Table C-1. ")
    r_cap_label.font.name = "Times New Roman"
    r_cap_label.font.size = Pt(10.5)
    r_cap_label.font.bold = True
    
    r_cap_text = p_cap.add_run("Summary Checklist of All User Manual Figures, Document Sections, Target Roles, and Functional Components")
    r_cap_text.font.name = "Times New Roman"
    r_cap_text.font.size = Pt(10.5)
    r_cap_text.font.italic = True
    
    # Data Rows Definition
    data = [
        (
            "Figure C-1",
            "The Public Website (No Account Needed)\nPublic Welcome Portal & Landing Gateway",
            "Public /\nCommuters /\nUnregistered Guests",
            "Santa Rosa Homes TODA brand identity & municipal mission banner; standardized municipal fare preview matrix; central terminal operational hours & dispatch policies; interactive FAQ accordion; \"Get Started\" commuter onboarding CTA; and secure \"Sign In\" portal trigger link."
        ),
        (
            "Figure C-2",
            "Account Access & Authentication\nSecure Portal Sign-In & Authentication Form",
            "All System Roles\n(Commuters, TODA Drivers, Admins, Superadmins)",
            "Role-aware registered email address and masked password input fields; interactive eye toggle for password visibility; \"Keep Me Signed In\" persistent session checkbox; primary \"Sign In\" submission CTA; \"Forgot Password?\" recovery trigger; and direct account registration link."
        ),
        (
            "Figure C-3",
            "Multi-Step Registration & Driver Document Submission\nPassenger Registration Flow (3 Panels)",
            "Commuters /\nPassengers",
            "3-step registration wizard with dynamic progress indicator bar; account role selection cards (Commuter vs. Driver); personal contact inputs (Full Legal Name, unique Mobile Number, Barangay, and verified Email); real-time duplicate validation badges; strong password creation with complexity validation rules; Terms of Service & Privacy Policy consent; and complete registration CTA button."
        ),
        (
            "Figure C-4",
            "Multi-Step Registration & Driver Document Submission\nDriver Registration Flow & 2FA OTP (2 Panels)",
            "TODA Drivers /\nApplicants",
            "Municipal MTOP franchise body number input; LGU Tricycle Permit and Professional Driver's License image upload dropzones with format/size validation rules; Two-Factor Authentication (2FA) Gmail OTP verification modal; 6-digit segmented input cells; 60-second OTP resend cooldown timer; and automated submission into the admin compliance review queue."
        ),
        (
            "Figure C-5",
            "Multi-Step Registration & Driver Document Submission\nDriver Accreditation Review Statuses (3 Panels)",
            "TODA Drivers /\nApplicants",
            "Driver accreditation lifecycle state deck: \"Under Review\" status banner with 24–48 hour verification SLA notice; \"Action Required\" defect alert card with administrative rejection notes and immediate document re-upload triggers; \"Disqualified\" ineligibility notice with station hotline link; and return-to-guest portal button."
        ),
        (
            "Figure C-6",
            "Passenger Experience & Commuter Hub\nPassenger Ride Booking Sheet & Fare Calculation Modal (2 Panels)",
            "Commuters /\nPassengers",
            "MapLibre GL vector map viewport with real-time GPS user geolocation pin; central TODA terminal station marker & 35m geofence boundary; available queued driver real-time counter pill; quick-select popular landmark carousel; pickup & drop-off destination inputs with interactive \"Pin on Map\" picker; passenger capacity stepper (1–4 riders); automated transparent fare computation based on road distance; and \"Confirm & Request Tricycle\" dispatch CTA to #1 queued driver."
        ),
        (
            "Figure C-7",
            "Passenger Experience & Commuter Hub\nFare Bargaining & Counter-Proposal Modals (2 Panels)",
            "TODA Drivers &\nCommuters",
            "Regulated fare negotiation workflow: 20-second dynamic decision countdown timer; pickup & destination trip summary card; driver fare adjuster stepper (₱5 increments); quick fare preset chips (₱20, ₱25, ₱30, ₱35, ₱40); \"Send Fare Proposal\" trigger; passenger-facing counter-offer display with price comparison; and 1-tap \"Accept Proposed Fare\" or \"Decline / Next Driver\" action triggers."
        ),
        (
            "Figure C-8",
            "Passenger Experience & Commuter Hub\nActive Ride Navigation & Real-Time In-App Trip Chat Drawer (2 Panels)",
            "Commuters &\nTODA Drivers",
            "Live transit execution interface: Assigned driver identity card (photo, verified name, 6-digit MTOP body number); pickup address card; cellular phone call trigger; slide-over in-app messaging drawer with unread message counter badge; quick-response message preset chips (\"On my way, 2 mins!\", \"I am waiting outside\", \"Thank you!\"); timestamped chronological message feed with read receipts; and 250-character live text input bar with instant WebSocket transmission."
        ),
        (
            "Figure C-9",
            "Passenger Experience & Commuter Hub\nDriver Arrival, Passenger Boarding & Drop-Off Completion (2 Panels)",
            "TODA Drivers",
            "Driver transit progression lifecycle: \"Arrived at Pickup\" status banner with automated audio chime and push notification dispatch; passenger boarding verification card with destination route overview and agreed fare; \"Start Trip / Depart\" trigger transitioning status to \"In Transit\"; live vector approach route line on map; drop-off establishment & passenger headcount summary; and \"Complete Drop Off\" CTA crediting driver wallet and activating terminal return guidance."
        ),
        (
            "Figure C-10",
            "Driver Operations & Queue Dispatch Hub\nTerminal Return Guidance & Passenger 5-Star Service Rating Modal (2 Panels)",
            "TODA Drivers &\nCommuters",
            "Post-trip completion workflows: Driver side receives real-time earnings credit toast (+₱60.00), turn-by-turn map route guiding tricycle back to Santa Rosa Homes central terminal, 35-meter geofence-restricted auto check-in trigger, and \"Wayside Passenger\" pickup button; Commuter side receives trip completed celebration banner, total fare paid receipt summary with driver MTOP info, interactive 1-to-5 star service rating selector with verbal descriptors, commendation chips (\"Safe Driving\", \"Polite Driver\", \"Clean Tricycle\", \"On-Time Arrival\"), and permanent rating submission CTA."
        ),
        (
            "Figure C-11",
            "Driver Operations & Queue Dispatch Hub\nDriver Queue Hero Card, Duty Toggle, & Active Queue Roster (2 Panels)",
            "TODA Drivers",
            "Terminal queue management console: Offline standby duty status indicator; floating \"Go On-Duty\" power switch; radiant green \"#1 in Queue\" hero card designating top dispatch priority; \"Next for Dispatch\" readiness notice; \"Start Terminal Walk-In Ride\" action button; live FIFO-ordered terminal queue roster of stationed drivers; active \"Drivers On Trip\" monitor; and dispatcher manual drag-and-drop queue reorder gripper."
        ),
        (
            "Figure C-12",
            "Driver Operations & Queue Dispatch Hub\nTerminal Walk-In & Returning Wayside Dispatch Modals (2 Panels)",
            "TODA Drivers",
            "Dual-channel non-app dispatching forms: Terminal Walk-In Dispatch modal with fixed station origin, destination barangay selection chips, passenger capacity stepper, and computed regulated tariff calculation; Wayside Dispatch modal enabling returning en route drivers to log flagged passengers, select drop-off zones, compute transit fare, update driver ledger, and prevent unauthorized unregistered trips."
        ),
        (
            "Figure C-13",
            "Trip History & Certified Digital Receipts\nTrip History Ledger & Certified Digital Fare Receipt Modal (2 Panels)",
            "All System Roles\n(Commuters &\nTODA Drivers)",
            "Auditable transit record interface: Destination search bar; date horizon filter chips (All, Today, This Week); completed trip summary cards with timestamps and fare badges; dispatch mode indicator (Online App Booking vs. Terminal Walk-In); certified digital fare receipt modal with unique alphanumeric trip reference ID; color-coded transit route timeline (origin to destination); verified driver & passenger identification stamp with MTOP badge; and itemized fare breakdown (base tariff, distance surcharge, total collected)."
        ),
        (
            "Figure C-14",
            "Driver Performance & Analytics\nDriver Rating Deck, 7-Day Revenue Trends, & Channel Share Analytics (2 Panels)",
            "TODA Drivers",
            "Driver performance & financial analytics dashboard: Timeframe filter pill switcher (Today, This Week, This Month, All Time); animated radial satisfaction rating ring displaying cumulative star score (e.g. 4.0 / 5.0); total completed trips counter; aggregate gross revenue KPI card; interactive 7-day revenue trend pillar chart highlighting peak earning day badge; and dispatch channel share breakdown bars comparing earnings across Terminal Queue, Online Dispatch, and Wayside Flagged rides."
        ),
        (
            "Figure C-15",
            "TODA Administration Console & Dispatch Management\nTODA Admin Console — Applicant Review Roster & Document Verification Modal (2 Panels)",
            "TODA Station Admins\n& Officers",
            "Station administration onboarding tool: Segmented console navigation bar (Drivers, Queue, Reports, Broadcast); applicant status filter chips (All, Pending Review, Approved, Suspended); applicant summary roster cards (name, submitted MTOP permit number, timestamp); digital verification modal header; high-resolution scanned MTOP and driver's license document viewer; administrative compliance notes textarea; and one-click \"Approve & Grant Accreditation\" or \"Reject Application\" decision buttons."
        ),
        (
            "Figure C-16",
            "TODA Administration Console & Dispatch Management\nTerminal Queue Drag-and-Drop Management & Incident Dispute Ledger (2 Panels)",
            "TODA Station Admins\n& Dispatchers",
            "Operational fleet supervision console: Live terminal queue line roster of stationed drivers; \"Emergency Reset Stalled Trip\" trigger to restore stalled driver units into queue sequence; commuter incident grievance dispute ledger with search bar and status filters (All, Pending, Investigating, Resolved); detailed passenger incident report cards (reference ID, violation category, complainant details, driver MTOP permit); and case adjudication status tags (Investigating, Resolved, Dismissed)."
        ),
        (
            "Figure C-17",
            "TODA Administration Console & Dispatch Management\nTODA Broadcast Announcement Dispatcher & Live Member Notification Drawer (2 Panels)",
            "TODA Admins &\nAssociation Members",
            "Association-wide communication hub: \"Create Announcement\" modal trigger for drafting meeting notices, fare advisories, or urgent weather warnings; published announcement bulletins archive with delivery counts; target audience filter tags (All Members, Drivers Only, Commuters Only); member slide-over notification drawer with unread counter badge; high-priority bulletin alert cards with issuance timestamps; and \"Mark All as Read & Dismiss\" batch acknowledgment controls."
        ),
        (
            "Figure C-18",
            "Official Report Generator & A4 Print Engine\nOfficial Report Generator — Template Configuration & Certified A4 PDF Preview (2 Panels)",
            "TODA Admins &\nCertifying Officers",
            "Municipal-grade reporting engine: 4 certified report template selection cards (Incident Matrix, Grouped Fleet Roster, Chronological Ledger, Operations Dashboard); date horizon filter pills (All Time, Today, This Week, Month-to-Date, Last Month, Custom Date Range); certifying officer signatory name and designation inputs; \"Generate Report\" primary CTA; interactive A4 sheet canvas with zoom controls (Fit to Page, 50% to 200%) and panning; and \"Print / Save PDF\" action button triggering standard A4 browser print engine with formal CSS print styling."
        ),
        (
            "Figure C-19",
            "Superadmin Master Governance Console\nSuperAdmin Governance Console — Executive Pulse & Fleet Operations",
            "Superadmins &\nLGU Regulators",
            "Executive governance overview: Master sidebar navigation with 7 core system modules (Overview, Users, Tariffs & Radar, MTOP Fleet, Bylaws, Branding, Audit Logs); real-time Executive Pulse server health banner (PHP 8.3.33 runtime, database storage size, operational uptime); key performance metric cards (gross dispatched volume, active registered fleet count); interactive 7-day dispatch volume SVG area chart plotting completed rides, walk-ins, and revenue; and real-time live audit ledger feed tracking active and completed dispatches."
        ),
        (
            "Figure C-20",
            "Superadmin Master Governance Console\nSuperAdmin Tariff Matrix Calibration & Terminal Geofence Radar Controls",
            "Superadmins &\nLGU Regulators",
            "Municipal fare & geospatial governance module: Regulated distance tariff input form configuring Base Fare (₱15.00), Succeeding Distance Rate (₱3.50/km), Night Differential Surcharge (₱5.00), and Terminal Fund Fee (₱2.00); central terminal GPS anchor coordinate inputs and check-in radar radius calibration (35 meters); master database table of popular community landmarks with fixed municipal fares; and \"Save Tariff & Radar Calibration\" CTA persisting updates across all commuter and driver mobile clients instantly."
        ),
        (
            "Figure C-21",
            "User Profile & Device Settings\nUser Profile Management, Accredited MTOP Badge, & Hardware Preferences (2 Panels)",
            "All System Roles\n(Commuters, TODA Drivers, Admins, Superadmins)",
            "Account profile & device diagnostics suite: User avatar photo container with floating camera upload FAB supporting real-time photo replacement; verified legal name display with official accredited MTOP franchise badge; daily trips count and earnings balance metric tiles; official compliance status indicator badge (Approved, Under Review, Suspended); hardware sensor permission status indicators (high-precision GPS geolocation, push notifications, and Web Audio API synthesized dispatch chime toggles); automated \"Run Device Sensor Diagnostics\" action tool; and secure \"Log Out & End Session\" termination button."
        )
    ]
    
    # Create Table: cols = 4, rows = len(data) + 1
    table = doc.add_table(rows=len(data) + 1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="B0BEC5", sz="4", val="single")
    
    col_widths = [Inches(1.0), Inches(2.0), Inches(1.3), Inches(2.2)]
    
    # Format Header Row
    hdr_cells = table.rows[0].cells
    hdr_titles = [
        "Figure Number",
        "Document Section & Form Name",
        "Primary Target Role",
        "Key Components & Functions Covered"
    ]
    
    make_row_header(table.rows[0])
    make_row_cant_split(table.rows[0])
    
    for c_i, title in enumerate(hdr_titles):
        cell = hdr_cells[c_i]
        cell.width = col_widths[c_i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, "1E293B")  # Dark slate navy
        set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    # Populate Data Rows
    for r_i, (fig_num, sec_form, role, components) in enumerate(data):
        row = table.rows[r_i + 1]
        make_row_cant_split(row)
        cells = row.cells
        
        # Zebra striping
        bg_color = "F8FAFC" if r_i % 2 == 1 else "FFFFFF"
        
        for c_i in range(4):
            cells[c_i].width = col_widths[c_i]
            cells[c_i].vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_background(cells[c_i], bg_color)
            set_cell_margins(cells[c_i], top=100, bottom=100, left=110, right=110)
            
        # Col 0: Figure Number (Centered, Bold)
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.15
        r0 = p0.add_run(fig_num)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
        # Col 1: Document Section & Form Name
        p1 = cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.15
        
        parts = sec_form.split("\n")
        # Section name (normal/italic)
        r1_sec = p1.add_run(parts[0] + "\n")
        r1_sec.font.name = "Times New Roman"
        r1_sec.font.size = Pt(9)
        r1_sec.font.italic = True
        r1_sec.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        # Form name (bold)
        r1_form = p1.add_run(parts[1] if len(parts) > 1 else "")
        r1_form.font.name = "Times New Roman"
        r1_form.font.size = Pt(9.5)
        r1_form.font.bold = True
        r1_form.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
        # Col 2: Primary Target Role (Centered / Left)
        p2 = cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(role)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(9.5)
        r2.font.bold = False
        r2.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        # Col 3: Key Components & Functions Covered (Justified)
        p3 = cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p3.paragraph_format.space_before = Pt(2)
        p3.paragraph_format.space_after = Pt(2)
        p3.paragraph_format.line_spacing = 1.15
        r3 = p3.add_run(components)
        r3.font.name = "Times New Roman"
        r3.font.size = Pt(9)
        r3.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
    # Concluding Note
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(14)
    p_note.paragraph_format.space_after = Pt(4)
    p_note.paragraph_format.line_spacing = 1.15
    p_note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    r_note_bold = p_note.add_run("Summary Overview: ")
    r_note_bold.font.name = "Times New Roman"
    r_note_bold.font.size = Pt(9.5)
    r_note_bold.font.bold = True
    r_note_bold.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    r_note_text = p_note.add_run(
        "Across all 21 user manual figures, the SRH LINK-TODA system establishes a comprehensive end-to-end digital dispatch, queue governance, tariff computation, communication, and reporting workflow tailored for the Santa Rosa Homes Tricycle Operators and Drivers Association. The architecture strictly isolates public fare inquiries from private queue telemetry and administrative compliance data while providing responsive vector mapping, automated SMS/Gmail 2FA security, dynamic fare bargaining, and certified A4 reporting standards."
    )
    r_note_text.font.name = "Times New Roman"
    r_note_text.font.size = Pt(9.5)
    r_note_text.font.italic = True
    r_note_text.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    try:
        doc.save(output_path)
        print(f"Successfully generated docx at: {output_path}")
    except PermissionError:
        alt_path = output_path.replace(".docx", "_New.docx")
        doc.save(alt_path)
        print(f"Note: '{output_path}' is currently open in Word. Saved copy to '{alt_path}' instead.")

if __name__ == "__main__":
    out1 = r"c:\laragon\www\srh-toda-app\User_Manual_Summary_Checklist.docx"
    out2 = r"c:\laragon\www\srh-toda-app\paper\User_Manual_Summary_Checklist.docx"
    out3 = r"c:\laragon\www\srh-toda-app\paper\User_Manual_Notes_and_Checklist.docx"
    build_summary_checklist_doc(out1)
    build_summary_checklist_doc(out2)
    build_summary_checklist_doc(out3)
