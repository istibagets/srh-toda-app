import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

FIGURES_DIR = r"paper\arranged screenshots"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    insideV = OxmlElement('w:insideV')
    insideV.set(qn('w:val'), 'none')
    tblBorders.append(insideV)
    tblPr.append(tblBorders)

def add_figure_image(doc, filename, max_width_inches=6.2):
    file_path = os.path.join(FIGURES_DIR, filename)
    if not os.path.exists(file_path):
        print(f"WARNING: Image not found: {file_path}")
        return
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run()
    run.add_picture(file_path, width=Inches(max_width_inches))

def add_table_descriptions(doc, rows_data):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    hdr_cells[0].width = Inches(0.8)
    hdr_cells[1].width = Inches(5.7)
    
    set_cell_background(hdr_cells[0], "F1F5F9")
    set_cell_background(hdr_cells[1], "F1F5F9")
    set_cell_margins(hdr_cells[0], top=100, bottom=100, left=120, right=120)
    set_cell_margins(hdr_cells[1], top=100, bottom=100, left=120, right=120)
    
    p0 = hdr_cells[0].paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run("No.")
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(10.5)
    r0.font.bold = True
    
    p1 = hdr_cells[1].paragraphs[0]
    r1 = p1.add_run("Field, Control, or Area — Detailed Technical Description")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    
    # Data rows
    for i, (num_str, title_str, desc_str) in enumerate(rows_data):
        row_cells = tbl.rows[i + 1].cells
        row_cells[0].width = Inches(0.8)
        row_cells[1].width = Inches(5.7)
        row_cells[0].vertical_alignment = WD_ALIGN_VERTICAL.TOP
        row_cells[1].vertical_alignment = WD_ALIGN_VERTICAL.TOP
        
        # Zebra striping
        if i % 2 == 1:
            set_cell_background(row_cells[0], "FAFAFA")
            set_cell_background(row_cells[1], "FAFAFA")
        else:
            set_cell_background(row_cells[0], "FFFFFF")
            set_cell_background(row_cells[1], "FFFFFF")
            
        set_cell_margins(row_cells[0], top=80, bottom=80, left=120, right=120)
        set_cell_margins(row_cells[1], top=80, bottom=80, left=120, right=120)
        
        # Col 0: Number
        cp0 = row_cells[0].paragraphs[0]
        cp0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp0.paragraph_format.space_before = Pt(2)
        cp0.paragraph_format.space_after = Pt(2)
        cr0 = cp0.add_run(num_str)
        cr0.font.name = 'Times New Roman'
        cr0.font.size = Pt(10)
        cr0.font.bold = True
        
        # Col 1: Title & Description
        cp1 = row_cells[1].paragraphs[0]
        cp1.paragraph_format.space_before = Pt(2)
        cp1.paragraph_format.space_after = Pt(2)
        cp1.paragraph_format.line_spacing = 1.15
        
        rtitle = cp1.add_run(f"{title_str} — ")
        rtitle.font.name = 'Times New Roman'
        rtitle.font.size = Pt(10)
        rtitle.font.bold = True
        
        rdesc = cp1.add_run(desc_str)
        rdesc.font.name = 'Times New Roman'
        rdesc.font.size = Pt(10)
        rdesc.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_figure_label(doc, fig_num, fig_title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    r_fig = p.add_run(f"Figure {fig_num}. ")
    r_fig.font.name = 'Times New Roman'
    r_fig.font.size = Pt(11)
    r_fig.font.bold = True
    
    r_title = p.add_run(fig_title)
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(11)
    r_title.font.italic = True
    return p

def main():
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.header.is_linked_to_previous = False
        
    # Document Title Block
    p_app = doc.add_paragraph()
    p_app.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_app.paragraph_format.space_before = Pt(0)
    p_app.paragraph_format.space_after = Pt(4)
    r_app = p_app.add_run("APPENDIX C")
    r_app.font.name = 'Times New Roman'
    r_app.font.size = Pt(14)
    r_app.font.bold = True
    
    p_man = doc.add_paragraph()
    p_man.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_man.paragraph_format.space_before = Pt(0)
    p_man.paragraph_format.space_after = Pt(16)
    r_man = p_man.add_run("User Manual: SRH LINK-TODA Queue Management System")
    r_man.font.name = 'Times New Roman'
    r_man.font.size = Pt(13)
    r_man.font.bold = True
    
    # =========================================================================
    # SECTION 1: PUBLIC WEBSITE
    # =========================================================================
    add_heading_1(doc, "The Public Website (No Account Needed)")
    add_heading_2(doc, "What the Landing Page Is For")
    add_body_p(doc, "The landing page serves as the public-facing entry point of the SRH LINK-TODA progressive web application. It introduces the purpose of the platform, presents standardized municipal fare rates, provides access to Santa Rosa Homes terminal operational information, answers frequently asked questions, and serves as the secure access gateway for passengers and accredited drivers.")
    add_body_p(doc, "An important feature of the public website is its privacy boundary. Confidential queue positions, active trip routes, driver mobile numbers, passenger identities, and administrative dispute records are not displayed publicly. Publicly presented information is strictly intended to explain system policies and fare transparency without exposing private transit information.")
    
    add_figure_label(doc, "C-1", "Public Welcome Portal & Landing Gateway")
    add_figure_image(doc, "Figure_C-01_Public_Landing_Page.png", max_width_inches=5.2)
    
    add_table_descriptions(doc, [
        ("1", "Brand Logo & Association Emblem", "Official motorized tricycle icon embodying the Santa Rosa Homes TODA community transit brand identity."),
        ("2", "Main System Title", "Prominent system headline identifying the official SRH LINK-TODA progressive web platform."),
        ("3", "Value Proposition Subtitle", "Introductory mission statement highlighting passenger safety, reliable scheduling, and verified local drivers."),
        ("4", "Get Started Primary CTA Button", "Primary action trigger routing new commuters and drivers directly into the multi-step account registration flow."),
        ("5", "Existing User Sign In Link", "Direct navigation button routing registered passengers, drivers, and administrators into the authentication portal.")
    ])

    # =========================================================================
    # SECTION 2: ACCESS & AUTHENTICATION
    # =========================================================================
    add_heading_1(doc, "Account Access & Authentication")
    add_heading_2(doc, "Signing In and Driver Two-Factor OTP Verification")
    add_body_p(doc, "Registered users access the system through the unified Sign-In page. Users enter their registered email address and password to log in. For TODA Driver accounts, the system enforces Two-Factor Gmail OTP Verification to safeguard driver queue credentials and earnings data against unauthorized device access.")
    
    add_figure_label(doc, "C-2", "Secure Portal Sign-In and Authentication Form")
    add_figure_image(doc, "Figure_C-02_Portal_Sign_In_Form.png", max_width_inches=5.2)
    
    add_table_descriptions(doc, [
        ("1", "Back Navigation Trigger", "Returns the user to the public welcome landing page."),
        ("2", "Welcome Back Header", "Contextual headline confirming the secure sign-in portal state."),
        ("3", "Email Address Input Field", "Validated text field equipped with email icon and one-touch clear button for entering account credentials."),
        ("4", "Password Input with Visibility Toggle", "Encrypted password input equipped with an interactive eye icon allowing users to inspect or mask entered characters."),
        ("5", "Keep Me Signed In Checkbox", "Enables persistent local session caching on personal mobile devices."),
        ("6", "Sign In Action Button", "Authenticates entered credentials against the backend API and initiates role-based dashboard routing."),
        ("7", "Create One Registration Link", "Direct route for commuters and drivers who have not yet established an account in the system.")
    ])

    add_heading_2(doc, "Multi-Step Registration & Driver Document Submission")
    add_body_p(doc, "The registration portal provides a structured multi-step wizard for onboarding new commuters and TODA drivers. It validates email and phone number uniqueness in real-time, enforces strong passwords, and collects driver MTOP franchise permits and driver's licenses for administrative accreditation.")
    
    add_figure_label(doc, "C-3", "Passenger Registration Flow — Role Selection, Personal Details, and Security Setup")
    add_figure_image(doc, "Figure_C-03_Passenger_Registration_Flow.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Registration Progress Indicator", "Visual 3-step progress bar tracking active setup percentage (33% to 100%)."),
        ("2", "Passenger Commuter Role Card", "Selection card activating commuter mode for booking tricycle rides across Santa Rosa Homes."),
        ("3", "TODA Driver Role Card", "Selection card activating driver mode requiring MTOP franchise validation and admin review."),
        ("4", "Step Continuation Button", "Validates active form constraints and transitions smoothly to the next pane."),
        ("5", "Full Name Input Field", "Captures complete legal user name with active validity checkmark."),
        ("6", "Email Address Field with Live Availability", "Real-time asynchronous checking ensuring the entered email address is unique and valid."),
        ("7", "Mobile Phone Number with Live Availability", "11-digit mobile phone field verified against existing records to prevent duplicate accounts."),
        ("8", "Password Strength Input", "Enforces robust password criteria (minimum 8 characters with alphanumeric requirements)."),
        ("9", "Confirm Password Validation Field", "Real-time character comparison ensuring error-free password entry."),
        ("10", "Complete Registration Action Button", "Finalizes commuter account creation and grants immediate access to the Commuter Hub.")
    ])

    add_figure_label(doc, "C-4", "Driver Registration Flow — MTOP Credentials, Document Upload & Two-Factor OTP")
    add_figure_image(doc, "Figure_C-04_Driver_Registration_Credentials.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Driver Accreditation Requirement Alert", "Informs applicants that valid MTOP franchise permits and LTO licenses are mandatory for review."),
        ("2", "Full Name on Driver's License", "Captures the exact legal name matching the official government driver's license."),
        ("3", "Official 6-Digit MTOP Body Number", "Restricted numeric input field recording the municipal franchise permit number."),
        ("4", "MTOP Certificate Dropzone", "Interactive file uploader accepting camera photos or PDF scans of the official MTOP franchise certificate."),
        ("5", "Driver's License Dropzone", "File uploader accepting scanned copies of the valid Land Transportation Office (LTO) license."),
        ("6", "Gmail OTP Security Notice", "Notifies the driver that a 6-digit cryptographic verification code has been dispatched to their email."),
        ("7", "6-Digit OTP Code Input Box", "Specialized numeric input field with automated formatting and character limit enforcement."),
        ("8", "Resend Verification Timer", "Enforces a 60-second cooldown period before a replacement OTP token can be dispatched."),
        ("9", "Verify & Complete Registration CTA", "Validates the OTP token against server cache and submits the application for administrative review.")
    ])

    add_figure_label(doc, "C-5", "Driver Accreditation Review Statuses — Pending, Action Required, and Disqualified")
    add_figure_image(doc, "Figure_C-05_Driver_Accreditation_Statuses.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Profile Under Review Header", "Confirms that driver registration documents have been received and queued for admin inspection."),
        ("2", "Submitted MTOP Permit Details", "Displays the applicant's recorded MTOP permit number undergoing verification."),
        ("3", "Verification Timeframe Notice", "Advises the applicant that document review typically concludes within 24 to 48 business hours."),
        ("4", "Application Needs Review Alert", "Alert notification indicating that submitted documentation failed compliance verification."),
        ("5", "Official Rejection Note & Compliance Feedback", "Displays specific remarks written by the TODA administrator detailing missing or illegible files."),
        ("6", "Submit Appeal / Re-upload CTA", "Opens the remediation form allowing the driver to submit updated credentials."),
        ("7", "Permanent Disqualification Notice", "Indicates that the driver account has been permanently revoked due to critical compliance violations."),
        ("8", "TODA Grievance Office Contact Link", "Provides administrative office location and telephone contact for formal in-person appeals.")
    ])

    # =========================================================================
    # SECTION 3: PASSENGER EXPERIENCE
    # =========================================================================
    add_heading_1(doc, "Passenger Experience & Commuter Hub")
    add_heading_2(doc, "Interactive Map Viewport & Commuter Action Sheet")
    add_body_p(doc, "The Passenger Commute Hub provides an intuitive vector-based map interface powered by MapLibre GL. Commuters can view their exact GPS location, observe available tricycle queue density at the central terminal, browse popular destinations with fixed fares, and open the ride request sheet.")
    
    add_figure_label(doc, "C-6", "Passenger Ride Booking Sheet and Interactive Fare Calculation Modal")
    add_figure_image(doc, "Figure_C-06_Passenger_Ride_Booking_Modal.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "MapLibre GL Vector Map Viewport", "Hardware-accelerated vector map rendering Santa Rosa Homes street layout, GPS positioning, and terminal hubs."),
        ("2", "Available Queue Count Pill", "Displays the real-time count of verified drivers currently waiting in queue at the TODA terminal station."),
        ("3", "Popular Landmarks Carousel", "Horizontal scroller of frequent destinations pre-configured with standardized municipal fares for 1-tap booking."),
        ("4", "Book a Tricycle Primary CTA", "Action button expanding the full pickup address and destination calculation modal."),
        ("5", "Pickup Address Field with GPS Auto-Fill", "Text box pre-populated with device GPS coordinates, editable for Block and Lot accuracy."),
        ("6", "Drop-Off Destination & Pin on Map Trigger", "Destination input field equipped with interactive map pin-drop mode for exact coordinate selection."),
        ("7", "Passenger Capacity Stepper", "Adjusts rider count (1 to 4 passengers) with automated tariff computation for excess capacity."),
        ("8", "Barangay Regulated Fare Breakdown Card", "Displays transparent computed tariff based on surveyed road distance and passenger volume."),
        ("9", "Confirm & Request Tricycle CTA", "Dispatches ride request directly to the prioritized #1 driver in the active terminal queue.")
    ])

    add_heading_2(doc, "Fare Bargaining, Driver Negotiation & Counter-Offers")
    add_body_p(doc, "The SRH LINK-TODA system integrates a regulated Fare Bargaining module. When special circumstances arise (e.g. heavy cargo, inclement weather, or specialized drop-offs), drivers can counter-propose adjusted fares using a ₱5-increment stepper, which passengers can review and accept or decline in real time.")
    
    add_figure_label(doc, "C-7", "Fare Bargaining and Counter-Proposal Negotiation Modals")
    add_figure_image(doc, "Figure_C-07_Fare_Bargaining_Negotiation.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "20-Second Decision Countdown Timer", "Dynamic progress countdown bar before the pending ride offer automatically rotates to the next driver in line."),
        ("2", "Pickup Location & Destination Landmark", "Summary card displaying commuter pickup address and destination spot."),
        ("3", "Fare Proposal Adjuster Stepper", "Interactive '+' and '-' stepper controls allowing drivers to adjust proposed rates in ₱5 increments."),
        ("4", "Quick Fare Preset Chips", "One-touch preset buttons (₱20, ₱25, ₱30, ₱35, ₱40) for rapid rate adjustment."),
        ("5", "Send Fare Proposal Action Button", "Transmits proposed rate to the passenger's screen for real-time confirmation."),
        ("6", "Driver Proposed Fare Alert Banner", "Passenger notification displaying the counter-proposed rate with transparent cost comparison."),
        ("7", "Accept Proposed Fare Action CTA", "Passenger button confirming the proposed rate and transitioning the ride into En Route state."),
        ("8", "Decline / Next Driver Button", "Declines counter-proposal and returns the passenger to regular dispatch queue matching.")
    ])

    add_heading_2(doc, "Active Ride Tracking & Real-Time In-App Trip Chat")
    add_body_p(doc, "Once a driver accepts the trip, the passenger sheet transitions through real-time lifecycle states (En Route, Arrived, In Transit). Commuters can track driver proximity, initiate cellular phone calls, exchange in-app chat messages, and access safety reporting features.")
    
    add_figure_label(doc, "C-8", "Active Ride Navigation and Real-Time In-App Trip Chat Drawer")
    add_figure_image(doc, "Figure_C-08_En_Route_Navigation_Trip_Chat.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Live En Route Route Line on Vector Map", "Real-time vector path illustrating assigned driver location and anticipated approach route."),
        ("2", "Passenger Identity Card & Pickup Address", "Displays commuter name, Block/Lot pickup details, and agreed ride tariff."),
        ("3", "One-Touch Cellular Phone Call Button", "Direct native dialer trigger establishing voice communication with the driver."),
        ("4", "Open Trip Chat Action Trigger", "Launches the slide-over real-time messaging drawer with unread message counter badge."),
        ("5", "Arrived at Pickup Transition Trigger", "Driver button confirming arrival outside the commuter's pickup location."),
        ("6", "Recipient Driver Identity & MTOP Header", "Displays assigned driver photo, verified full name, and 6-digit MTOP body number."),
        ("7", "Quick Response Preset Chips", "One-tap preset chips ('On my way, 2 mins!', 'I am waiting outside', 'Thank you!') for rapid hands-free messaging."),
        ("8", "Chronological Message Feed", "Encrypted real-time message stream with timestamped speech bubbles and read receipts."),
        ("9", "Message Input Bar & Send Trigger", "Text input field supporting up to 250 characters with instant WebSocket transmission.")
    ])

    add_figure_label(doc, "C-9", "Driver Arrival, Passenger Boarding & Drop-Off Completion")
    add_figure_image(doc, "Figure_C-09_Driver_Arrival_Dropoff_Completion.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Arrived at Pickup Status Banner", "Alert bar confirming vehicle arrival, notifying commuter via audio chime and push notification."),
        ("2", "Passenger Boarding & Route Overview", "Displays verified passenger details, route destination, and agreed fare collection amount."),
        ("3", "Start Trip / Depart Action Button", "Driver trigger confirming passenger has safely boarded, transitioning ride status to 'In Transit'."),
        ("4", "On Trip In Transit Status Indicator", "Active transit state banner tracking live travel towards drop-off destination."),
        ("5", "Drop-Off Landmark & Passenger Count", "Destination overview card displaying drop-off establishment and passenger headcount."),
        ("6", "Complete Drop Off Action Button", "Finalizes the ride upon safe arrival, credits fare to driver wallet, and activates terminal return workflow.")
    ])

    # =========================================================================
    # SECTION 4: DRIVER OPERATIONS
    # =========================================================================
    add_heading_1(doc, "Driver Operations & Queue Dispatch Hub")
    add_heading_2(doc, "Duty Status, Geofenced Queue Positioning & Hero Card")
    add_body_p(doc, "The Driver Queue Hub manages operator duty state, geofenced terminal check-in, queue priority order, and active trip execution. When drivers enter the 35-meter terminal geofence, toggling On Duty automatically registers them into the First-In, First-Out (FIFO) queue.")
    
    add_figure_label(doc, "C-10", "Terminal Return Guidance and Passenger 5-Star Service Rating Modal")
    add_figure_image(doc, "Figure_C-10_Terminal_Return_Passenger_Rating.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Drop-Off Earnings Credit Toast", "Real-time toast notification confirming fare credited to driver balance (+₱60.00)."),
        ("2", "Terminal Station Return Guidance", "Map route guiding driver back to the central Santa Rosa Homes terminal station."),
        ("3", "Terminal Queue Check-In Button", "Geofence-restricted trigger allowing driver to re-enter the queue upon arriving within 35 meters."),
        ("4", "Trip Completed Celebration Header", "Passenger confirmation banner verifying arrival at destination."),
        ("5", "Total Fare Paid & Driver MTOP Info", "Itemized receipt summary displaying total payment and assigned driver MTOP number."),
        ("6", "Interactive 5-Star Service Rating Selector", "Allows passengers to rate service quality from 1 to 5 stars with dynamic verbal descriptors."),
        ("7", "Commendation Chips", "Selectable compliment tags ('Safe Driving', 'Polite Driver', 'Clean Tricycle', 'On-Time Arrival')."),
        ("8", "Submit Rating & Feedback CTA", "Saves rating to driver's permanent record and updates community safety ratings.")
    ])

    add_figure_label(doc, "C-11", "Driver Queue Hero Card, Duty Toggle, and Active Queue Roster")
    add_figure_image(doc, "Figure_C-11_Driver_Queue_Hero_Card.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Floating Go On-Duty Power Switch", "Interactive toggle button switching driver between On-Duty (Active Queue) and Off-Duty states."),
        ("2", "Offline Duty Status Indicator", "Displays current driver standby state ('YOU'RE OFFLINE') and prompt to activate queue participation."),
        ("3", "Geofence Terminal Check-In Guidance", "Contextual instruction advising driver to enter within 35 meters of the terminal station."),
        ("4", "Radiant Green #1 in Queue Hero Card", "Dynamic animated hero deck highlighting top queue priority for terminal and app dispatches."),
        ("5", "Next for Dispatch Status Notice", "Official dispatch notification alerting driver to prepare vehicle for upcoming passenger matching."),
        ("6", "Start Terminal Walk-In Ride Button", "Action button allowing prioritized driver to log passengers boarding directly at the station."),
        ("7", "Active Terminal Queue Roster", "Live list of all accredited drivers currently queued at the terminal, ordered by FIFO timestamp.")
    ])

    add_figure_label(doc, "C-12", "Terminal Walk-In and Returning Wayside Dispatch Modals")
    add_figure_image(doc, "Figure_C-12_WalkIn_Wayside_Dispatch.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Terminal Dispatch Modal Header", "Header identifying station-origin walk-in dispatch mode."),
        ("2", "Pickup Origin Point", "Fixed origin anchor set to the central TODA station terminal."),
        ("3", "Destination Barangay Selection Chips", "Quick-select grid of popular barangays and zones with surveyed distance metrics."),
        ("4", "Passenger Stepper & Computed Regulated Tariff", "Rider volume selector computing official tariff based on municipal rate matrices."),
        ("5", "Depart & Start Walk-In Trip Button", "Dispatches unit, records trip in audit ledger, and temporarily removes driver from queue."),
        ("6", "Wayside Dispatch Modal Header", "Header identifying transit-origin wayside dispatch mode for flagging passengers en route."),
        ("7", "Dynamic En Route GPS Origin Address", "Automatically resolves driver's current coordinates along the roadway as trip origin."),
        ("8", "Wayside Destination Selector & Fare Card", "Destination selector calculating return fare and passenger volume surcharge."),
        ("9", "Depart & Start Wayside Trip Button", "Launches wayside trip and tracks transit progress back toward terminal.")
    ])

    # =========================================================================
    # SECTION 5: HISTORY & RECEIPTS
    # =========================================================================
    add_heading_1(doc, "Trip History & Certified Digital Receipts")
    add_heading_2(doc, "Trip Ledger & Digital Receipt Modal")
    add_body_p(doc, "The History Ledger maintains an auditable record of all completed, walk-in, and cancelled trips for both passengers and drivers. Users can search by destination, filter by date range, and view itemized digital receipts.")
    
    add_figure_label(doc, "C-13", "Trip History Ledger and Certified Digital Fare Receipt Modal")
    add_figure_image(doc, "Figure_C-13_Trip_History_Digital_Receipt.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Search Bar & Time Horizon Chips", "Real-time search filtering trips by destination, passenger, or trip ID, with timeframe chips (All, Today, Week)."),
        ("2", "Completed Trip Summary Card", "Card displaying unique trip reference ID, timestamp, pickup/drop-off nodes, and fare paid."),
        ("3", "Dispatch Mode Badge", "Identifies whether the trip was an Online App Booking or Terminal Walk-In dispatch."),
        ("4", "Official Digital Fare Receipt Header", "Modal header displaying certified trip ID number and official completion timestamp."),
        ("5", "Transit Route Timeline", "Color-coded route node breakdown detailing origin address and destination drop-off location."),
        ("6", "Participants & Vehicle Verification Stamp", "Displays assigned driver full name, official MTOP franchise badge, and passenger details."),
        ("7", "Itemized Fare Breakdown & Total Collected", "Itemized breakdown covering base tariff, distance surcharge, and gross fare collected.")
    ])

    # =========================================================================
    # SECTION 6: DRIVER PERFORMANCE & EARNINGS
    # =========================================================================
    add_heading_1(doc, "Driver Performance & Analytics")
    add_heading_2(doc, "Executive Rating Deck, 7-Day Revenue Chart & Channel Share")
    add_body_p(doc, "The Driver Performance Module visualizes satisfaction ratings via an animated radial ring, breaks down weekly revenue trends across daily pillar charts, and illustrates dispatch channel proportions.")
    
    add_figure_label(doc, "C-14", "Driver Rating Deck, 7-Day Revenue Trends, and Channel Share Analytics")
    add_figure_image(doc, "Figure_C-14_Driver_Earnings_Performance_Analytics.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Time Horizon Filter Switcher", "Pill selector toggling performance statistics across Today, This Week, This Month, and All Time."),
        ("2", "Radial Satisfaction Rating Ring", "Conic progress indicator displaying cumulative passenger service satisfaction (e.g. 4.0 / 5.0 stars)."),
        ("3", "Total Completed Trips Done Counter", "Metric tile tracking total successful dispatches executed within the selected period."),
        ("4", "Aggregate Gross Revenue KPI Tile", "Financial card tracking gross revenue accumulated across all executed trips."),
        ("5", "7-Day Revenue Trend Pillar Chart", "Interactive bar chart illustrating day-by-day revenue fluctuations and peak demand days."),
        ("6", "Peak Earning Day Analytics Badge", "Identifies the highest grossing day of the reporting cycle to aid schedule optimization."),
        ("7", "Dispatch Channel Share Breakdown", "Visual proportion bars tracking revenue generated via Terminal Queue, Online Dispatch, and Wayside rides.")
    ])

    # =========================================================================
    # SECTION 7: TODA ADMINISTRATION CONSOLE
    # =========================================================================
    add_heading_1(doc, "TODA Administration Console & Dispatch Management")
    add_heading_2(doc, "Driver Accreditation, Live Queue Reordering & Dispute Resolution")
    add_body_p(doc, "The TODA Admin Console empowers station officers to inspect uploaded MTOP permits and driver's licenses, manage driver compliance status, reorder active queue priority via drag-and-drop, resolve commuter incident reports, and broadcast association bulletins.")
    
    add_figure_label(doc, "C-15", "TODA Admin Console — Applicant Review Roster and Document Verification Modal")
    add_figure_image(doc, "Figure_C-15_Admin_Driver_Accreditation_Review.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Admin Segmented Navigation Bar", "Navigation tabs switching between Drivers, Queue, Reports, and Broadcast administrative consoles."),
        ("2", "Applicant Status Filter Chips", "Filter chips segregating driver applicants across All, Pending Review, Approved, and Suspended."),
        ("3", "Driver Applicant Summary Card", "Roster card displaying applicant name, submitted MTOP permit number, and submission timestamp."),
        ("4", "Digital Verification Modal Header", "Official modal interface for inspecting uploaded franchise certificates and government licenses."),
        ("5", "Scanned MTOP Document High-Res Viewer", "Interactive document viewer rendering high-resolution uploaded certificate images and PDFs."),
        ("6", "Administrative Compliance Notes Textarea", "Official text area for recording inspection findings, compliance defects, or approval justifications."),
        ("7", "Approve & Grant Accreditation Action Button", "Grants formal accreditation, enabling driver to check into the terminal queue and accept rides."),
        ("8", "Reject Application Action Button", "Declines applicant with recorded reasons, transmitting instant notification to the driver.")
    ])

    add_figure_label(doc, "C-16", "Terminal Queue Drag-and-Drop Management and Incident Dispute Ledger")
    add_figure_image(doc, "Figure_C-16_Admin_Terminal_Queue_Dispute_Ledger.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Active Terminal Queue Line Roster", "Live list of drivers stationed at the terminal awaiting dispatch matching."),
        ("2", "Drag-and-Drop Queue Reorder Gripper", "Interactive handle allowing dispatchers to manually adjust terminal queue order during operational anomalies."),
        ("3", "Emergency Reset Stalled Trip Button", "Resets stalled or abandoned trips, returning driver unit to terminal queue sequence."),
        ("4", "Incident Search Bar & Status Filter Horizon", "Real-time search and status filter tabs (All, Pending, Investigating, Resolved) for commuter complaints."),
        ("5", "Passenger Incident Report Card", "Detailed grievance card detailing report reference ID, incident category, passenger name, and driver MTOP."),
        ("6", "Case Adjudication Status Tag", "Displays handling state (Investigating, Resolved, Dismissed) with link to view full case logs.")
    ])

    add_figure_label(doc, "C-17", "TODA Broadcast Announcement Dispatcher and Live Member Notification Drawer")
    add_figure_image(doc, "Figure_C-17_Admin_Broadcast_Announcements.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Create Announcement Trigger Button", "Opens broadcast composer modal for drafting general assembly notices or urgent advisories."),
        ("2", "Published Announcement Bulletins Stream", "Chronological archive of past association announcements with read confirmation counts."),
        ("3", "Target Audience Filter Tag", "Identifies message audience filter (All Members, Drivers Only, or Commuters Only)."),
        ("4", "Notifications Drawer Header & Counter", "Member notification slide-over header displaying count of unread community announcements."),
        ("5", "Live Official Bulletin Alert Card", "High-priority alert card detailing broadcast title, issuance timestamp, and message body."),
        ("6", "Mark All as Read & Dismiss Controls", "Quick action button marking all received announcements as acknowledged.")
    ])

    # =========================================================================
    # SECTION 8: OFFICIAL REPORT GENERATOR
    # =========================================================================
    add_heading_1(doc, "Official Report Generator & A4 Print Engine")
    add_heading_2(doc, "Template Configuration & Certified A4 PDF Preview")
    add_body_p(doc, "The Official Report Generator produces certified executive reports for TODA meetings, Barangay audits, and municipal compliance. It supports 4 certified formats with real-time A4 sheet rendering, zoom controls, and browser PDF printing.")
    
    add_figure_label(doc, "C-18", "Official Report Generator — Template Configuration and Live Certified A4 Document Preview")
    add_figure_image(doc, "Figure_C-18_Admin_Official_Report_Generator.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Report Template Selection Cards", "Tile deck selecting among 4 certified templates (Incident Matrix, Grouped Fleet Roster, Chronological Ledger, Operations Dashboard)."),
        ("2", "Date Horizon Filter Horizon Pills", "Filters data timeframe (All Time, Today, This Week, Month-to-Date, Last Month, Custom Date Range)."),
        ("3", "Signatory Name & Designation Inputs", "Configures official certifying officer name and association designation printed on the document footer."),
        ("4", "Generate Report Primary CTA", "Compiles selected database records and renders the high-fidelity A4 document canvas."),
        ("5", "Interactive A4 Sheet Zoom Controls", "Controls canvas scaling (Fit to Page, 50% to 200%) and interactive panning for document inspection."),
        ("6", "Official Republic of the Philippines TODA Letterhead", "Formal header containing association logo, municipal seal, and Barangay office address."),
        ("7", "Certified Tabular Audit Matrix", "High-contrast audit data table structured for official municipal compliance and Barangay records."),
        ("8", "Formal Signatory Certification Block", "Printed signature line, officer designation, generation timestamp, and system verification seal."),
        ("9", "Print / Save PDF Action Button", "Launches browser print engine with standard A4 margins and formal CSS print media rules.")
    ])

    # =========================================================================
    # SECTION 9: SUPERADMIN MASTER CONSOLE
    # =========================================================================
    add_heading_1(doc, "Superadmin Master Governance Console")
    add_heading_2(doc, "Executive Pulse Dashboard & System Health")
    add_body_p(doc, "The Superadmin Governance Console provides executive oversight over system health, dispatch volume trends, revenue distribution, user accounts, and immutable audit logs.")
    
    add_figure_label(doc, "C-19", "SuperAdmin Governance Console — Executive Pulse and Fleet Operations")
    add_figure_image(doc, "Figure_C-19_SuperAdmin_Pulse_MTOP_Registry.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Master Governance Sidebar Navigation", "Sidebar granting immediate access to all 7 core modules (Overview, Users, Tariffs & Radar, MTOP Fleet, Bylaws, Branding, Audit Logs)."),
        ("2", "Executive Pulse Server Health Banner", "Real-time system health monitor displaying server environment (PHP 8.3.33), database size, and operational uptime."),
        ("3", "System Key Performance Metric Cards", "KPI deck tracking gross dispatched volume, HOA terminal maintenance fund share (₱2.00/trip), and active fleet count."),
        ("4", "Dispatch Volume 7-Day Area Chart", "Visual SVG chart plotting daily completed trips, walk-in dispatches, and accumulated revenue."),
        ("5", "Recent Dispatches Live Audit Ledger Feed", "Real-time audit stream displaying active and recently completed trips, assigned drivers, MTOP permits, and collected fares.")
    ])

    add_heading_2(doc, "Tariff Calibration, Geofence Radar & Boundary Controls")
    add_body_p(doc, "Administrators can calibrate municipal distance tariffs, set the terminal geofence coordinates and radar radius (35m), configure preset destination landmarks, and manage user account privileges.")
    
    add_figure_label(doc, "C-20", "SuperAdmin Tariff Matrix Calibration and Terminal Geofence Radar Controls")
    add_figure_image(doc, "Figure_C-20_SuperAdmin_Tariff_Radar_Security.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Municipal Distance Tariff Inputs", "Form fields configuring Base Fare (₱15.00), Succeeding Distance Rate (₱3.50/km), Night Differential (₱5.00), and Terminal Fund Fee (₱2.00)."),
        ("2", "Terminal Geofence GPS Coordinates & Radius", "Configures exact terminal latitude/longitude anchor and check-in radar radius (35 meters)."),
        ("3", "Popular TODA Landmarks & Fixed Municipal Fares Table", "Master table of pre-configured community destinations with standardized fixed fares."),
        ("4", "Save Tariff & Radar Calibration CTA", "Persists updated fare schedules and geofence rules immediately across all client applications."),
        ("5", "Interactive Terminal Geofence Radar Visualizer", "Vector visualizer mapping terminal perimeter boundaries and active driver positioning radar.")
    ])

    # =========================================================================
    # SECTION 10: USER PROFILE & DEVICE SETTINGS
    # =========================================================================
    add_heading_1(doc, "User Profile & Device Settings")
    add_heading_2(doc, "Profile Management & Hardware Sensor Diagnostics")
    add_body_p(doc, "The Profile Module enables users across all roles to manage their personal contact details, upload profile photos, inspect franchise certificates (for drivers), update passwords, and configure hardware sensor preferences (GPS, push notifications, chime audio).")
    
    add_figure_label(doc, "C-21", "User Profile Management, Accredited MTOP Badge, and Hardware Preferences")
    add_figure_image(doc, "Figure_C-21_Profile_Management_Permissions.png", max_width_inches=6.2)
    
    add_table_descriptions(doc, [
        ("1", "Profile Avatar Photo Frame & Camera Upload FAB", "User avatar container with floating camera button supporting gallery upload and real-time photo replacement."),
        ("2", "Driver Name & Accredited MTOP Franchise Badge", "Displays verified legal name and official municipal MTOP franchise body number badge."),
        ("3", "Daily Trips & Earnings Performance Metric Tiles", "Tiles tracking daily completed trips and accumulated earnings balance."),
        ("4", "Compliance Status Indicator Badge", "Displays official accreditation state (Approved, Under Review, or Suspended)."),
        ("5", "GPS High-Accuracy Location Permission Toggle", "Hardware permission switch enabling high-precision GPS geolocation required for terminal check-in."),
        ("6", "Push Notifications Toggle Switch", "Interactive toggle enabling native push notification alerts for ride requests and system updates."),
        ("7", "In-App Chime Audio Synthesizer Toggle Switch", "Toggle switch activating audible alert sound effects and chimes for queue updates."),
        ("8", "Run Device Sensor Diagnostics Action Button", "Automated diagnostic tool evaluating browser GPS accuracy, notification service workers, and Web Audio API synthesizer.")
    ])

    # =========================================================================
    # SUMMARY CHECKLIST TABLE
    # =========================================================================
    add_heading_1(doc, "Summary Checklist of All User Manual Figures")
    add_body_p(doc, "The following master index summarizes all twenty-one (21) user manual figures documented in Appendix C, including their target user roles and primary functional components:")
    
    summary_data = [
        ("Figure C-1", "Public Welcome Portal & Landing Gateway", "Public / Commuters", "Brand logo, municipal mission, Get Started CTA, Sign In trigger"),
        ("Figure C-2", "Secure Portal Sign-In & Authentication Form", "All Users / Drivers", "Email/Password fields, eye visibility toggle, keep signed in, sign in CTA"),
        ("Figure C-3", "Passenger Registration Flow (3 Panels)", "Commuters", "Role selection, personal details, live validation, security password"),
        ("Figure C-4", "Driver Registration Flow & 2FA OTP (2 Panels)", "TODA Drivers", "MTOP body number, franchise dropzones, Gmail OTP verification"),
        ("Figure C-5", "Driver Accreditation Review Statuses (3 Panels)", "TODA Drivers", "Pending review, action required feedback, disqualified notice"),
        ("Figure C-6", "Passenger Ride Booking Sheet & Modal (2 Panels)", "Commuters", "MapLibre GL map, available driver count, landmark carousel, fare card"),
        ("Figure C-7", "Fare Bargaining & Counter-Proposal (2 Panels)", "Drivers / Commuters", "20s decision timer, ₱5 fare stepper, preset chips, accept/decline"),
        ("Figure C-8", "En Route Navigation & Trip Chat (2 Panels)", "Drivers / Commuters", "Active vector route, phone call, preset quick chips, live chat feed"),
        ("Figure C-9", "Driver Arrival & Drop-Off Completion (2 Panels)", "Drivers", "Arrived alert, boarding status, destination summary, complete drop-off"),
        ("Figure C-10", "Terminal Return & 5-Star Rating (2 Panels)", "Drivers / Commuters", "Earnings toast, terminal return route, 5-star selector, compliment tags"),
        ("Figure C-11", "Driver Queue Hero Card & Duty Toggle (2 Panels)", "TODA Drivers", "Offline switch, Go On Duty toggle, #1 in Queue hero card, active queue"),
        ("Figure C-12", "Terminal Walk-In & Wayside Dispatch (2 Panels)", "TODA Drivers", "Station origin, wayside en route GPS, destination chips, regulated tariff"),
        ("Figure C-13", "Trip History Ledger & Digital Receipt (2 Panels)", "All Users", "Search bar, date horizon filters, trip ledger, certified fare receipt"),
        ("Figure C-14", "Driver Earnings & Performance Analytics (2 Panels)", "TODA Drivers", "Radial rating ring, trips counter, 7-day revenue chart, channel share"),
        ("Figure C-15", "Admin Driver Accreditation Review (2 Panels)", "TODA Admins", "Applicant roster, document inspection modal, approve/reject triggers"),
        ("Figure C-16", "Terminal Queue Reordering & Disputes (2 Panels)", "TODA Admins", "Drag-to-reorder queue line, reset trip, incident dispute ledger"),
        ("Figure C-17", "Admin Broadcast & Member Notifications (2 Panels)", "TODA Admins / Members", "Create announcement, published archive, slide-over notification drawer"),
        ("Figure C-18", "Official Report Generator & A4 Print Engine (2 Panels)", "TODA Admins / Officers", "4 report templates, timeframe filters, certified A4 printable sheet"),
        ("Figure C-19", "SuperAdmin Governance Console Dashboard", "Superadmins", "Executive pulse server monitor, KPI metric tiles, 7-day area chart"),
        ("Figure C-20", "SuperAdmin Tariff Calibration & Geofence Radar", "Superadmins", "Distance rate matrix, 35m geofence coordinates, popular landmarks table"),
        ("Figure C-21", "User Profile Management & App Permissions (2 Panels)", "All Users", "Avatar upload, MTOP badge, sensor permissions, device diagnostics")
    ]
    
    tbl_sum = doc.add_table(rows=len(summary_data) + 1, cols=4)
    tbl_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sum)
    
    col_widths = [Inches(1.0), Inches(1.8), Inches(1.5), Inches(2.2)]
    
    hdr_sum = tbl_sum.rows[0].cells
    hdr_titles = ["Figure Number", "Document Section & Form Name", "Primary Target Role", "Key Components Covered"]
    for j, (hcell, htitle) in enumerate(zip(hdr_sum, hdr_titles)):
        hcell.width = col_widths[j]
        set_cell_background(hcell, "F1F5F9")
        set_cell_margins(hcell, top=100, bottom=100, left=100, right=100)
        p = hcell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(htitle)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        
    for i, (fig_no, form_name, role_name, comps) in enumerate(summary_data):
        row_cells = tbl_sum.rows[i + 1].cells
        for j, (cell, val) in enumerate(zip(row_cells, [fig_no, form_name, role_name, comps])):
            cell.width = col_widths[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if i % 2 == 1:
                set_cell_background(cell, "FAFAFA")
            else:
                set_cell_background(cell, "FFFFFF")
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.5)
            if j == 0 or j == 1:
                r.font.bold = True

    # Output paths
    candidate_paths = [
        r"c:\laragon\www\srh-toda-app\paper\Appendix_C_User_Manual.docx",
        r"c:\laragon\www\srh-toda-app\paper\Appendix_C_User_Manual_Updated.docx",
        r"c:\laragon\www\srh-toda-app\paper\Appendix_C_User_Manual_Clean.docx",
        r"c:\laragon\www\srh-toda-app\paper\Appendix_C_User_Manual_With_Callouts.docx",
        r"c:\laragon\www\srh-toda-app\Appendix_C_User_Manual.docx",
        r"c:\laragon\www\srh-toda-app\Appendix_C_User_Manual_Clean.docx",
    ]
    
    saved_paths = []
    for p in candidate_paths:
        try:
            doc.save(p)
            print(f"Successfully saved: {p}")
            saved_paths.append(p)
        except PermissionError:
            print(f"Could not write to {p} (locked by Word).")
        except Exception as e:
            print(f"Failed to write {p}: {e}")
            
    if saved_paths:
        print(f"\nDocument generation complete! Generated {len(saved_paths)} copies.")
    else:
        print("\nWarning: All candidate paths were locked.")

if __name__ == "__main__":
    main()
